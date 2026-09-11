#!/usr/bin/env python3
"""Read-only preview of the Fantasy History workbook import.

This script does not write app data. It reads the raw matchup rows from the
archived workbook and prints the facts we expect to preserve.
"""

from __future__ import annotations

import argparse
import json
import sys
from collections import defaultdict
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ModuleNotFoundError:  # pragma: no cover - local environment guidance
    print(
        "Missing Python package: openpyxl. Run this through Codex, or install "
        "openpyxl in your local Python environment.",
        file=sys.stderr,
    )
    raise


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = PROJECT_ROOT / "reference" / "workbooks" / "Fantasy History.xlsx"

RAW_COLUMNS = {
    "year": 9,
    "week": 10,
    "team1": 11,
    "team2": 12,
    "team1_score": 13,
    "team2_score": 14,
    "regular": 15,
    "playoff": 16,
    "consolation": 17,
    "final_seeding": 18,
}

EXPECTED_SEASONS = {
    2015: {"rows": 102, "weeks": (1, 17), "teams": 12, "regular": 78, "playoff": 6, "consolation": 18, "final_seeding": 1},
    2016: {"rows": 95, "weeks": (1, 16), "teams": 12, "regular": 78, "playoff": 5, "consolation": 12, "final_seeding": 1},
    2017: {"rows": 95, "weeks": (1, 16), "teams": 12, "regular": 78, "playoff": 5, "consolation": 12, "final_seeding": 1},
    2018: {"rows": 95, "weeks": (1, 16), "teams": 12, "regular": 78, "playoff": 5, "consolation": 12, "final_seeding": 1},
    2019: {"rows": 95, "weeks": (1, 16), "teams": 12, "regular": 78, "playoff": 5, "consolation": 12, "final_seeding": 1},
    2020: {"rows": 95, "weeks": (1, 16), "teams": 12, "regular": 78, "playoff": 5, "consolation": 12, "final_seeding": 1},
    2021: {"rows": 101, "weeks": (1, 17), "teams": 12, "regular": 84, "playoff": 5, "consolation": 12, "final_seeding": 1},
    2022: {"rows": 101, "weeks": (1, 17), "teams": 12, "regular": 84, "playoff": 5, "consolation": 12, "final_seeding": 1},
}


@dataclass(frozen=True)
class RawMatchupPreview:
    row_number: int
    season: int
    week: int
    team1: str
    team2: str
    team1_score: float
    team2_score: float
    game_type: str
    final_seeding_game: bool
    result: str
    margin: float


def cell_value(row: tuple[Any, ...], column_number: int) -> Any:
    return row[column_number - 1] if len(row) >= column_number else None


def enabled(value: Any) -> bool:
    return value in (1, "1", True)


def to_int(value: Any, label: str, row_number: int, issues: list[str]) -> int | None:
    if isinstance(value, bool) or value is None:
        issues.append(f"Row {row_number}: missing or invalid {label}.")
        return None

    try:
        numeric_value = int(value)
    except (TypeError, ValueError):
        issues.append(f"Row {row_number}: {label} is not a whole number: {value!r}.")
        return None

    if numeric_value != value and not isinstance(value, str):
        issues.append(f"Row {row_number}: {label} has an unexpected value: {value!r}.")
        return None

    return numeric_value


def to_score(value: Any, label: str, row_number: int, issues: list[str]) -> float | None:
    if isinstance(value, bool) or value is None:
        issues.append(f"Row {row_number}: missing or invalid {label}.")
        return None

    try:
        score = float(value)
    except (TypeError, ValueError):
        issues.append(f"Row {row_number}: {label} is not numeric: {value!r}.")
        return None

    return round(score, 2)


def get_game_type(row: tuple[Any, ...], row_number: int, issues: list[str]) -> str | None:
    primary_flags = {
        "regular": enabled(cell_value(row, RAW_COLUMNS["regular"])),
        "playoff": enabled(cell_value(row, RAW_COLUMNS["playoff"])),
        "consolation": enabled(cell_value(row, RAW_COLUMNS["consolation"])),
    }
    selected = [name for name, is_selected in primary_flags.items() if is_selected]

    if len(selected) != 1:
        issues.append(
            f"Row {row_number}: expected one primary game type, found {selected or 'none'}."
        )
        return None

    return selected[0]


def parse_workbook(workbook_path: Path) -> tuple[list[RawMatchupPreview], list[str]]:
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    if "Game Log" not in workbook.sheetnames:
        return [], ['Workbook is missing the "Game Log" sheet.']

    sheet = workbook["Game Log"]
    records: list[RawMatchupPreview] = []
    issues: list[str] = []

    for row_number, row in enumerate(sheet.iter_rows(min_row=4, values_only=True), start=4):
        raw_values = {name: cell_value(row, column) for name, column in RAW_COLUMNS.items()}

        if all(value is None for value in raw_values.values()):
            continue

        season = to_int(raw_values["year"], "season", row_number, issues)
        week = to_int(raw_values["week"], "week", row_number, issues)
        team1 = raw_values["team1"]
        team2 = raw_values["team2"]
        team1_score = to_score(raw_values["team1_score"], "Team 1 Score", row_number, issues)
        team2_score = to_score(raw_values["team2_score"], "Team 2 Score", row_number, issues)
        game_type = get_game_type(row, row_number, issues)

        if not isinstance(team1, str) or not team1.strip():
            issues.append(f"Row {row_number}: missing Team 1.")
            team1 = None
        if not isinstance(team2, str) or not team2.strip():
            issues.append(f"Row {row_number}: missing Team 2.")
            team2 = None
        if team1 == team2 and team1 is not None:
            issues.append(f"Row {row_number}: Team 1 and Team 2 are the same.")

        if None in (season, week, team1, team2, team1_score, team2_score, game_type):
            continue

        if team1_score > team2_score:
            result = f"{team1} win"
        elif team2_score > team1_score:
            result = f"{team2} win"
        else:
            result = "tie"

        records.append(
            RawMatchupPreview(
                row_number=row_number,
                season=season,
                week=week,
                team1=team1.strip(),
                team2=team2.strip(),
                team1_score=team1_score,
                team2_score=team2_score,
                game_type=game_type,
                final_seeding_game=enabled(raw_values["final_seeding"]),
                result=result,
                margin=round(abs(team1_score - team2_score), 2),
            )
        )

    return records, issues


def summarize(records: list[RawMatchupPreview], issues: list[str]) -> dict[str, Any]:
    managers = sorted({record.team1 for record in records} | {record.team2 for record in records})
    seasons: dict[int, dict[str, Any]] = defaultdict(
        lambda: {
            "rows": 0,
            "weeks": set(),
            "teams": set(),
            "regular": 0,
            "playoff": 0,
            "consolation": 0,
            "final_seeding": 0,
        }
    )

    for record in records:
        season = seasons[record.season]
        season["rows"] += 1
        season["weeks"].add(record.week)
        season["teams"].update([record.team1, record.team2])
        season[record.game_type] += 1
        if record.final_seeding_game:
            season["final_seeding"] += 1

    season_summary = {}
    for year, values in sorted(seasons.items()):
        weeks = sorted(values["weeks"])
        season_summary[year] = {
            "rows": values["rows"],
            "weeks": [weeks[0], weeks[-1]] if weeks else [],
            "teams": len(values["teams"]),
            "regular": values["regular"],
            "playoff": values["playoff"],
            "consolation": values["consolation"],
            "final_seeding": values["final_seeding"],
        }

    validations = validate_summary(records, managers, season_summary, issues)

    return {
        "matchup_rows": len(records),
        "seasons": sorted(season_summary),
        "manager_count": len(managers),
        "managers": managers,
        "season_summary": season_summary,
        "sample_rows": [asdict(record) for record in records[:5]],
        "validations": validations,
        "issues": issues,
    }


def validate_summary(
    records: list[RawMatchupPreview],
    managers: list[str],
    season_summary: dict[int, dict[str, Any]],
    issues: list[str],
) -> list[dict[str, str]]:
    validations: list[dict[str, str]] = []

    def add(name: str, passed: bool, detail: str) -> None:
        validations.append({"status": "OK" if passed else "CHECK", "name": name, "detail": detail})

    add("raw matchup rows", len(records) == 779, f"found {len(records)}, expected 779")
    add(
        "season range",
        sorted(season_summary) == list(EXPECTED_SEASONS),
        f"found {sorted(season_summary)}",
    )
    add("historical manager count", len(managers) == 16, f"found {len(managers)}, expected 16")
    add("parse issues", len(issues) == 0, f"found {len(issues)} issue(s)")

    for year, expected in EXPECTED_SEASONS.items():
        actual = season_summary.get(year)
        if not actual:
            add(f"{year} season present", False, "missing")
            continue

        for key in ("rows", "teams", "regular", "playoff", "consolation", "final_seeding"):
            add(
                f"{year} {key}",
                actual[key] == expected[key],
                f"found {actual[key]}, expected {expected[key]}",
            )

        add(
            f"{year} week span",
            tuple(actual["weeks"]) == expected["weeks"],
            f"found {actual['weeks']}, expected {list(expected['weeks'])}",
        )

    return validations


def print_human(summary: dict[str, Any], workbook_path: Path) -> None:
    print("Workbook import preview")
    print(f"Source: {workbook_path}")
    print()
    print(f"Raw matchup rows: {summary['matchup_rows']}")
    print(f"Seasons: {', '.join(str(season) for season in summary['seasons'])}")
    print(f"Managers: {summary['manager_count']}")
    print()

    print("Per-season shape")
    for year, values in summary["season_summary"].items():
        week_start, week_end = values["weeks"]
        print(
            f"- {year}: rows={values['rows']}, weeks={week_start}-{week_end}, "
            f"teams={values['teams']}, regular={values['regular']}, "
            f"playoff={values['playoff']}, consolation={values['consolation']}, "
            f"final_seeding={values['final_seeding']}"
        )
    print()

    print("Managers found")
    for manager in summary["managers"]:
        print(f"- {manager}")
    print()

    print("First five raw rows")
    for record in summary["sample_rows"]:
        print(
            f"- row {record['row_number']}: {record['season']} week {record['week']}, "
            f"{record['team1']} {record['team1_score']} vs "
            f"{record['team2']} {record['team2_score']} "
            f"({record['game_type']}, {record['result']})"
        )
    print()

    print("Validation")
    for validation in summary["validations"]:
        print(f"- {validation['status']}: {validation['name']} ({validation['detail']})")

    if summary["issues"]:
        print()
        print("Issues")
        for issue in summary["issues"][:25]:
            print(f"- {issue}")
        if len(summary["issues"]) > 25:
            print(f"- ...and {len(summary['issues']) - 25} more")


def main() -> int:
    parser = argparse.ArgumentParser(description="Preview Fantasy History workbook import.")
    parser.add_argument(
        "workbook",
        nargs="?",
        default=DEFAULT_WORKBOOK,
        type=Path,
        help="Path to the workbook to preview.",
    )
    parser.add_argument("--json", action="store_true", help="Print machine-readable JSON.")
    args = parser.parse_args()

    workbook_path = args.workbook.expanduser().resolve()
    if not workbook_path.exists():
        print(f"Workbook not found: {workbook_path}", file=sys.stderr)
        return 1

    records, issues = parse_workbook(workbook_path)
    summary = summarize(records, issues)

    if args.json:
        print(json.dumps(summary, indent=2, sort_keys=True))
    else:
        print_human(summary, workbook_path)

    has_failed_validation = any(item["status"] != "OK" for item in summary["validations"])
    return 1 if has_failed_validation else 0


if __name__ == "__main__":
    raise SystemExit(main())
